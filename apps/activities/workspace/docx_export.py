"""Deterministic DOCX rendering of the validated lecture; never executes model-generated code."""
from io import BytesIO
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .schemas import SECTION_IDS, validate_document

DOCX_MIME = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
MAX_DOCX_BYTES = 512_000


def filename(title):
    clean = re.sub(r'[^\w\s-]', '', title, flags=re.UNICODE).strip()[:140].strip()
    return (clean or 'Лекция') + '.docx'


def render_docx(lecture, section_ids=SECTION_IDS):
    validate_document(lecture, section_ids)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(2.2)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (('Title', 22), ('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 12)):
        style = document.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
    document.core_properties.title = lecture['title']
    document.core_properties.author = 'AI Edu'
    document.add_paragraph(lecture['title'], 'Title')
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    def inline(paragraph, nodes):
        for node in nodes:
            if node['type'] == 'hardBreak':
                paragraph.add_run().add_break()
                continue
            run = paragraph.add_run(node.get('text', ''))
            for mark in node.get('marks', []):
                if mark['type'] == 'bold':
                    run.bold = True
                elif mark['type'] == 'italic':
                    run.italic = True
                elif mark['type'] == 'code':
                    run.font.name, run.font.size = 'Courier New', Pt(10)
                elif mark['type'] == 'link':
                    # Keep links readable and inert, including when printed.
                    paragraph.add_run(' (' + mark['attrs']['href'] + ')')

    def numbering(kind, start):
        root = document.part.numbering_part.element
        abstract_id = max([int(el.get(qn('w:abstractNumId'))) for el in root.findall(qn('w:abstractNum'))] + [-1]) + 1
        abstract = OxmlElement('w:abstractNum')
        abstract.set(qn('w:abstractNumId'), str(abstract_id))
        lvl = OxmlElement('w:lvl')
        lvl.set(qn('w:ilvl'), '0')
        for tag, val in (('start', str(start)), ('numFmt', kind), ('lvlText', '•' if kind == 'bullet' else '%1.')):
            el = OxmlElement('w:' + tag)
            el.set(qn('w:val'), val)
            lvl.append(el)
        abstract.append(lvl)
        root.insert_element_before(abstract, 'w:num', 'w:numIdMac')
        num = root.add_num(abstract_id)
        return num.numId

    def blocks(nodes, depth=0):
        for node in nodes:
            kind = node['type']
            children = node.get('content', [])
            if kind in ('bulletList', 'orderedList'):
                num_id = numbering('bullet' if kind == 'bulletList' else {
                    '1': 'decimal', 'a': 'lowerLetter', 'A': 'upperLetter', 'i': 'lowerRoman', 'I': 'upperRoman'
                }.get(node.get('attrs', {}).get('type'), 'decimal'), node.get('attrs', {}).get('start', 1))
                for item in children:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.left_indent = Cm(.6 + depth * .5)
                    paragraph.paragraph_format.first_line_indent = Cm(-.4)
                    properties = paragraph._p.get_or_add_pPr()
                    num_pr = OxmlElement('w:numPr')
                    for tag, val in (('ilvl', '0'), ('numId', str(num_id))):
                        el = OxmlElement('w:' + tag)
                        el.set(qn('w:val'), val)
                        num_pr.append(el)
                    properties.append(num_pr)
                    inline(paragraph, item['content'][0].get('content', []))
                    blocks(item['content'][1:], depth + 1)
            elif kind == 'blockquote':
                before = len(document.paragraphs)
                blocks(children, depth)
                for paragraph in document.paragraphs[before:]:
                    paragraph.paragraph_format.left_indent = Cm(.8)
            elif kind in ('paragraph', 'heading', 'codeBlock'):
                style = f"Heading {node['attrs']['level']}" if kind == 'heading' else 'Normal'
                paragraph = document.add_paragraph(style=style)
                inline(paragraph, children)
                if kind == 'codeBlock':
                    for run in paragraph.runs:
                        run.font.name, run.font.size = 'Courier New', Pt(10)

    for part in lecture['sections']:
        document.add_paragraph(part['title'], 'Heading 1')
        blocks(part['content'].get('content', []))
    stream = BytesIO()
    document.save(stream)
    data = stream.getvalue()
    if len(data) > MAX_DOCX_BYTES:
        raise ValueError('docx_too_large')
    return data
