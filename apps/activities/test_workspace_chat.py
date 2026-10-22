import json
import uuid
from datetime import timedelta
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile
from docx import Document
from django.test import TestCase, SimpleTestCase, override_settings
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from apps.education.models import LearningItem, PracticeSubmission
from .models import WorkspaceArtifact, WorkspaceDocument, WorkspaceOperation, WorkspaceUsage
from .test_workspace import WorkspaceFixture, lecture, ai_response
from .workspace.docx_export import render_docx
from .workspace.limits import reserve_usage, ensure_capacity, WorkspaceLimit
from .workspace.schemas import validate_lecture
from .workspace.services import create_operation
from .workspace.worker import claim_next_operation, execute_operation


def plan(action='generate', **extra):
    return ai_response({'action': action, 'message': 'Готовлю лекцию.', 'choices': [],
                        'requirements': {'topic': 'Педагогика', 'audience': 'Студенты',
                                         'duration_minutes': 45, 'language': 'ru'}, 'docx': False, **extra})


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'],
                   AI_BASE_URL='https://provider.test/v1', AI_MODEL='test-model')
class WorkspaceChatTests(WorkspaceFixture, TestCase):
    def chat(self, responses, prompt='Подготовь лекцию'):
        self.document.refresh_from_db()
        data = self.operation_data('chat', prompt=prompt)
        data['request_id'] = uuid.UUID(data['request_id'])
        operation = create_operation(self.user, self.document, data)
        claim = claim_next_operation()
        self.assertEqual(claim.pk, operation.pk)
        with patch('apps.activities.workspace.conversation.run_ai', side_effect=responses) as provider:
            execute_operation(claim)
        operation.refresh_from_db()
        self.document.refresh_from_db()
        return operation, provider

    def history_operation(self):
        return WorkspaceOperation.objects.create(document=self.document, request_id=uuid.uuid4(), kind='chat',
            base_revision=self.document.revision, status='SUCCEEDED', input_data={}, model='test-model', timeout_seconds=120)

    def test_clarification_is_a_saved_message_without_generation_charge(self):
        operation, provider = self.chat([plan('clarify', message='Для какой аудитории?',
            choices=['Студенты', 'Преподаватели'], requirements={'topic': 'Педагогика'})])
        self.assertEqual(operation.status, 'SUCCEEDED')
        self.assertEqual(provider.call_count, 1)
        self.assertEqual(self.document.content, {})
        self.assertEqual(self.document.conversation_state['questions'], 1)
        self.assertEqual(self.document.conversation_state['requirements'], {'topic': 'Педагогика'})
        self.assertFalse(WorkspaceUsage.objects.exists())
        messages = self.client.get(self.path + 'messages/').data['results']
        self.assertEqual(messages[-1]['text'], 'Для какой аудитории?')
        self.assertEqual(messages[-1]['choices'], ['Студенты', 'Преподаватели'])

    def test_complete_request_generates_without_extra_questions(self):
        operation, provider = self.chat([plan(), ai_response()])
        self.assertEqual(operation.status, 'SUCCEEDED')
        self.assertEqual(self.document.content, lecture())
        self.assertEqual(WorkspaceUsage.objects.get().category, 'generate')
        self.assertEqual(operation.total_tokens, 600)
        self.assertEqual(len(operation.result['choices']), 5)
        self.assertEqual(provider.call_count, 2)

    def test_fourth_clarification_is_forced_to_generate(self):
        self.document.conversation_state = {'questions': 3, 'requirements': {'topic': 'Педагогика'}}
        self.document.save()
        operation, provider = self.chat([plan('clarify', message='Ещё один вопрос?'), ai_response()])
        self.assertEqual(operation.status, 'SUCCEEDED')
        self.assertEqual(operation.result['action'], 'generate')
        self.assertEqual(self.document.conversation_state['questions'], 0)
        self.assertEqual(provider.call_count, 2)

    def test_requirements_and_file_intent_survive_clarifications(self):
        self.chat([plan('clarify', message='Для кого?', choices=['Студенты'],
                        requirements={'topic': 'Педагогика'}, docx=True)])
        operation, provider = self.chat([plan('generate', requirements={'audience': 'Студенты'}, docx=True), ai_response()])
        self.assertEqual(operation.status, 'SUCCEEDED')
        self.assertEqual(self.document.parameters['topic'], 'Педагогика')
        self.assertEqual(self.document.parameters['audience'], 'Студенты')
        self.assertTrue(WorkspaceArtifact.objects.filter(operation=operation).exists())
        context = json.dumps(provider.call_args_list[0].args[0].input_data, ensure_ascii=False)
        self.assertIn('Для кого?', context)

    def test_edit_uses_current_lecture_and_separate_quota(self):
        self.document.content = lecture()
        self.document.save()
        changed = lecture()
        changed['sections'][5]['content']['content'][0]['content'][0]['text'] = 'Новый пример'
        operation, provider = self.chat([plan('edit'), ai_response(changed)], 'Добавь пример в раздел 6')
        self.assertEqual(operation.status, 'SUCCEEDED')
        self.assertEqual(self.document.content, changed)
        self.assertEqual(WorkspaceUsage.objects.get().category, 'edit')
        messages = provider.call_args_list[1].args[0].input_data['messages']
        self.assertIn(json.dumps(lecture(), ensure_ascii=False), [m['content'] for m in messages])

    def test_docx_download_is_private_immutable_and_free_to_repeat(self):
        operation, _ = self.chat([plan(docx=True), ai_response()])
        self.assertEqual(operation.status, 'SUCCEEDED')
        artifact = operation.result['artifact']
        original = self.client.get(artifact['url'])
        self.assertEqual(original.status_code, 200)
        self.assertTrue(original.content.startswith(b'PK'))
        self.assertIn('.docx', original['Content-Disposition'])
        count = WorkspaceUsage.objects.count()
        self.assertEqual(self.client.get(artifact['url']).content, original.content)
        self.assertEqual(WorkspaceUsage.objects.count(), count)
        self.document.content = {**lecture(), 'title': 'Другая версия'}
        self.document.save()
        self.assertEqual(self.client.get(artifact['url']).content, original.content)
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(artifact['url']).status_code, 404)
        self.client.logout()
        self.assertEqual(self.client.get(artifact['url']).status_code, 403)

    def test_export_existing_lecture_does_not_regenerate_content(self):
        self.document.content = lecture()
        self.document.save()
        operation, provider = self.chat([plan('export', docx=True)], 'Сделай Word-файл')
        self.assertEqual(operation.status, 'SUCCEEDED')
        self.assertEqual(provider.call_count, 1)
        self.assertEqual(list(WorkspaceUsage.objects.values_list('category', flat=True)), ['docx'])
        self.assertEqual(self.document.content, lecture())

    def test_export_without_existing_lecture_creates_full_lecture_and_file(self):
        route = plan('export')
        route['payload'].pop('docx')
        operation, provider = self.chat([route, ai_response()], 'Подготовь лекцию в Word')
        self.assertEqual(operation.status, 'SUCCEEDED')
        self.assertEqual(operation.result['action'], 'generate')
        self.assertEqual(self.document.content, lecture())
        self.assertTrue(operation.result.get('artifact'))
        self.assertEqual(provider.call_count, 2)

    def test_unsupported_formats_never_create_files(self):
        operation, provider = self.chat([plan('unsupported')], 'Сделай PPTX вместо лекции')
        self.assertEqual(operation.status, 'SUCCEEDED')
        self.assertEqual(operation.result['action'], 'unsupported')
        self.assertFalse(WorkspaceArtifact.objects.exists())
        self.assertFalse(WorkspaceUsage.objects.exists())
        self.assertEqual(provider.call_count, 1)

    def test_model_cannot_select_arbitrary_tools(self):
        operation, provider = self.chat([plan('execute_shell')], 'Запусти команду')
        self.assertEqual(operation.status, 'FAILED')
        self.assertEqual(operation.error_code, 'ai_invalid_response')
        self.assertFalse(WorkspaceArtifact.objects.exists())
        self.assertEqual(provider.call_count, 1)

    def test_hourly_generation_limit_stops_before_expensive_generation(self):
        for _ in range(5):
            reserve_usage(self.history_operation(), ['generate'])
        operation, provider = self.chat([plan()])
        self.assertEqual(operation.status, 'FAILED')
        self.assertEqual(operation.error_code, 'workspace_limit')
        self.assertEqual(provider.call_count, 1)
        self.assertEqual(self.document.content, {})

    def test_quota_reservations_are_atomic_idempotent_and_expire(self):
        for category, maximum in [('generate', 5), ('edit', 10), ('docx', 10)]:
            for _ in range(maximum):
                operation = self.history_operation()
                reserve_usage(operation, [category])
                reserve_usage(operation, [category])
            with self.assertRaises(WorkspaceLimit):
                reserve_usage(self.history_operation(), [category])
            self.assertEqual(WorkspaceUsage.objects.filter(category=category).count(), maximum)
        WorkspaceUsage.objects.filter(category='generate').update(reserved_at=timezone.now() - timedelta(hours=2))
        operation = self.history_operation()
        with self.assertRaises(WorkspaceLimit):
            reserve_usage(operation, ['generate', 'docx'])
        self.assertFalse(WorkspaceUsage.objects.filter(operation=operation).exists())
        reserve_usage(operation, ['generate'])

    def test_user_capacity_is_ten_across_documents(self):
        for index in range(10):
            item = LearningItem.objects.create(title=f'Other {index}', module=self.module, type='PRACTICE', position=index + 1)
            submission = PracticeSubmission.objects.create(user=self.user, item=item,
                                                           practice_version=self.document.submission.practice_version)
            document = WorkspaceDocument.objects.create(owner=self.user, submission=submission, document_type='lecture')
            WorkspaceOperation.objects.create(document=document, kind='chat', request_id=uuid.uuid4(), base_revision=1,
                input_data={}, model='test-model', timeout_seconds=120)
            if index < 9:
                ensure_capacity(self.user)
        with self.assertRaises(WorkspaceLimit):
            ensure_capacity(self.user)
        ensure_capacity(self.other)
        response = self.client.post(self.path + 'operations/', self.operation_data('chat'), format='json')
        self.assertEqual(response.status_code, 429)

    def test_oversized_lecture_is_not_saved(self):
        too_long = lecture()
        too_long['sections'][0]['content']['content'][0]['content'][0]['text'] = 'слово ' * 3001
        operation, _ = self.chat([plan(), ai_response(too_long)])
        self.assertEqual(operation.status, 'FAILED')
        self.assertEqual(operation.error_code, 'ai_invalid_content')
        self.assertEqual(self.document.content, {})
        self.assertFalse(WorkspaceArtifact.objects.exists())


class LectureDocxTests(SimpleTestCase):
    def test_docx_preserves_unicode_headings_marks_lists_and_page_numbers(self):
        content = lecture()
        content['title'] = 'Жасанды интеллект және білім беру'
        content['sections'][0]['content']['content'] = [
            {'type': 'heading', 'attrs': {'level': 2}, 'content': [{'type': 'text', 'text': 'Подраздел'}]},
            {'type': 'orderedList', 'attrs': {'start': 7}, 'content': [
                {'type': 'listItem', 'content': [{'type': 'paragraph', 'content': [
                    {'type': 'text', 'text': 'Важный пункт', 'marks': [{'type': 'bold'}]}]},
                    {'type': 'bulletList', 'content': [{'type': 'listItem', 'content': [
                        {'type': 'paragraph', 'content': [{'type': 'text', 'text': 'Вложенный пункт'}]}]}]}]}]},
            {'type': 'codeBlock', 'content': [{'type': 'text', 'text': 'print(1)'}]},
        ]
        data = render_docx(content)
        document = Document(BytesIO(data))
        self.assertEqual(document.paragraphs[0].text, content['title'])
        self.assertEqual(document.paragraphs[0].style.name, 'Title')
        self.assertEqual(sum(p.style.name == 'Heading 1' for p in document.paragraphs), 10)
        self.assertTrue(next(p for p in document.paragraphs if p.text == 'Важный пункт').runs[0].bold)
        with ZipFile(BytesIO(data)) as archive:
            numbering = archive.read('word/numbering.xml').decode()
            self.assertIn('w:val="7"', numbering)
            self.assertIn('PAGE', archive.read('word/footer1.xml').decode())
        self.assertLess(len(data), 512_000)

    def test_docx_and_chat_enforce_the_same_word_limit(self):
        content = lecture()
        content['sections'][0]['content']['content'][0]['content'][0]['text'] = 'слово ' * 3001
        with self.assertRaises(ValidationError):
            validate_lecture(content)
        with self.assertRaises(ValidationError):
            render_docx(content)
