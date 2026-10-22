import io
import json
from datetime import timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import Mock, patch
from types import SimpleNamespace
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TransactionTestCase, SimpleTestCase, override_settings
from django.utils import timezone
from rest_framework.test import APIClient
from apps.accounts.models import User
from apps.education.models import Course, Lesson, Module, LearningItem, LearningTrack
from apps.lessons.models import LectureImportJob
from apps.lessons.importing.document import to_blocks, ImportFailure
from apps.lessons.importing.worker import claim, execute, store_result


def office_file(kind='docx'):
    data = io.BytesIO()
    if kind == 'docx':
        from docx import Document
        document = Document()
        document.add_heading('Source section', 1)
        document.add_paragraph('Exact quotation: 42 != 24', style='Quote')
        document.add_paragraph('First item', style='List Bullet')
        document.add_paragraph('Second item', style='List Bullet')
        paragraph = document.add_paragraph()
        paragraph.add_run('Bold source').bold = True
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Value'
        table.cell(0, 1).text = '42'
    else:
        from pptx import Presentation
        from pptx.util import Inches
        document = Presentation()
        slide = document.slides.add_slide(document.slide_layouts[6])
        # XML order is intentionally opposite to visual reading order.
        slide.shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text = 'Below'
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Above'
    document.save(data)
    return SimpleUploadedFile(f'lecture.{kind}', data.getvalue())


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class PracticeInstructionImportTests(TransactionTestCase):
    def setUp(self):
        folder = TemporaryDirectory(); self.addCleanup(folder.cleanup)
        override = override_settings(MEDIA_ROOT=folder.name, LECTURE_IMPORT_TIMEOUT=90)
        override.enable(); self.addCleanup(override.disable)
        track = LearningTrack.objects.create(title='Import track', is_published=True)
        course = Course.objects.create(title='Import course', learning_track=track, is_published=True)
        self.module = Module.objects.create(course=course, title='Module', is_published=True)
        self.item = LearningItem.objects.create(module=self.module, title='Original task', type='PRACTICE')
        self.admin = User.objects.create_user(email='practice-import-admin@example.test', password='Password1234!', role='ADMIN')
        user = User.objects.create_user(email='practice-import-student@example.test', password='Password1234!', learning_track=track)
        self.client = APIClient(); self.client.force_authenticate(self.admin)
        self.student = APIClient(); self.student.force_authenticate(user)
        self.url = f'/api/v1/items/{self.item.short_id}/practice/imports/'
        self.draft_url = f'/api/v1/items/{self.item.short_id}/practice/draft/'

    def enqueue(self):
        response = self.client.post(self.url, {'file': office_file()}, format='multipart')
        self.assertEqual(response.status_code, 202, response.data)
        self.assertTrue(response.data['source_url'].startswith(self.url))
        return LectureImportJob.objects.get(pk=response.data['id'])

    def save_instructions(self, instructions):
        payload = {'practice_type': 'lecture_workshop', 'title': self.item.title, 'description': '',
                   'config': {'instructions': instructions, 'criteria': [{'title': 'Keep this criterion', 'description': '', 'is_required': True}]}}
        response = self.client.put(self.draft_url, payload, format='json')
        self.assertEqual(response.status_code, 200, response.data)
        return payload

    def test_real_import_requires_review_and_preserves_structure_when_published(self):
        job = self.enqueue()
        self.assertIsNone(job.lesson_id)
        self.assertEqual(job.practice_item_id, self.item.pk)
        execute(claim()); job.refresh_from_db()
        self.assertEqual(job.status, 'SUCCEEDED', job.error_code)
        self.item.refresh_from_db()
        self.assertIsNone(self.item.practice_id)
        instructions = [block['content'] for block in job.result['blocks']]
        payload = self.save_instructions(instructions)
        self.assertEqual(self.client.post(self.draft_url).status_code, 200)
        state = self.student.get(f'/api/v1/items/{self.item.short_id}/practice/').data
        self.assertEqual(state['version']['config'], payload['config'])
        self.assertIn('table', instructions[0])

    def test_history_sources_and_jobs_are_private_and_scoped_to_live_practices(self):
        job = self.enqueue()
        self.assertEqual(self.client.post(self.url, {'file': office_file()}, format='multipart').status_code, 409)
        self.assertEqual(self.student.get(self.url).status_code, 403)
        self.assertEqual(self.student.post(self.url, {'file': office_file()}, format='multipart').status_code, 403)
        self.assertEqual(self.student.get(self.url + f'{job.pk}/').status_code, 403)
        self.assertEqual(self.student.get(self.url + f'{job.pk}/source/').status_code, 403)
        source = self.client.get(self.url + f'{job.pk}/source/')
        self.assertEqual(source.status_code, 200); source.close()
        other = LearningItem.objects.create(module=self.module, title='Other', type='PRACTICE', position=1)
        self.assertEqual(self.client.get(f'/api/v1/items/{other.short_id}/practice/imports/{job.pk}/').status_code, 404)
        other.type = 'LECTURE'; other.save()
        self.assertEqual(self.client.get(f'/api/v1/items/{other.short_id}/practice/imports/').status_code, 404)
        self.item.delete()
        self.assertEqual(self.client.get(self.url + f'{job.pk}/').status_code, 404)

    def test_imported_images_become_accessible_only_after_publication(self):
        from PIL import Image
        self.enqueue(); job = claim()
        with TemporaryDirectory() as folder:
            Image.new('RGB', (8, 8), 'white').save(Path(folder) / 'figure.png')
            store_result(job, {'assets': [{'key': 'figure', 'file': 'figure.png', 'is_page': False}],
                              'elements': [{'source_id': 'e1', 'kind': 'image', 'image': 'figure', 'text': 'Task diagram'}], 'warnings': []}, folder)
        job.refresh_from_db(); media_id = job.result['blocks'][0]['media']
        media_url = f'/api/v1/media/{media_id}/'
        self.assertEqual(self.student.get(media_url).status_code, 403)
        instruction = json.dumps({'type': 'doc', 'content': [{'type': 'practiceImage', 'attrs': {'media': media_id, 'caption': 'Task diagram'}}]})
        self.save_instructions([instruction])
        self.assertEqual(self.student.get(media_url).status_code, 403)
        self.assertEqual(self.client.post(self.draft_url).status_code, 200)
        response = self.student.get(media_url)
        self.assertEqual(response.status_code, 200); response.close()
        call_command('cleanup_unused_media', delete=True, stdout=io.StringIO())
        response = self.student.get(media_url)
        self.assertEqual(response.status_code, 200); response.close()


class LectureImportTests(TransactionTestCase):
    def setUp(self):
        folder = TemporaryDirectory(); self.addCleanup(folder.cleanup)
        override = override_settings(MEDIA_ROOT=folder.name, LECTURE_IMPORT_TIMEOUT=90)
        override.enable(); self.addCleanup(override.disable)
        course = Course.objects.create(title='Import course')
        module = Module.objects.create(course=course, title='Module')
        self.lesson = Lesson.objects.create(module=module, title='Existing title')
        self.user = User.objects.create_user(email='import-admin@example.test', password='Password1234!', role='ADMIN')
        self.client = APIClient(); self.client.force_authenticate(self.user)
        student = User.objects.create_user(email='import-student@example.test', password='Password1234!')
        self.student = APIClient(); self.student.force_authenticate(student)
        self.url = f'/api/v1/lessons/{self.lesson.short_id}/imports/'
        self.draft_url = f'/api/v1/lessons/{self.lesson.short_id}/draft/'

    def enqueue(self, kind='docx'):
        response = self.client.post(self.url, {'file': office_file(kind)}, format='multipart')
        self.assertEqual(response.status_code, 202, response.data)
        return LectureImportJob.objects.get(pk=response.data['id'])

    def test_real_word_import_stays_private_and_preserves_structure_until_published(self):
        queued = self.enqueue()
        self.assertEqual(self.student.get(self.url).status_code, 403)
        self.assertEqual(self.student.get(self.url + f'{queued.pk}/source/').status_code, 403)
        source = self.client.get(self.url + f'{queued.pk}/source/')
        self.assertEqual(source.status_code, 200); source.close()
        execute(claim())
        queued.refresh_from_db()
        self.assertEqual(queued.status, 'SUCCEEDED', queued.error_code)
        block = queued.result['blocks'][0]
        nodes = json.loads(block['content'])['content']
        self.assertEqual([node['type'] for node in nodes], ['heading', 'blockquote', 'bulletList', 'paragraph', 'table'])
        self.assertEqual(nodes[1]['content'][0]['content'][0]['text'], 'Exact quotation: 42 != 24')
        self.assertEqual(nodes[3]['content'][0]['marks'], [{'type': 'bold'}])
        self.lesson.refresh_from_db()
        self.assertIsNone(self.lesson.draft_data)
        self.assertFalse(self.lesson.blocks.exists())
        snapshot = self.client.get(self.draft_url).data
        saved = self.client.put(self.draft_url, {**snapshot, 'blocks': [block]}, format='json')
        self.assertEqual(saved.status_code, 200, saved.data)
        self.assertFalse(self.lesson.blocks.exists())
        published = self.client.post(self.draft_url, {'revision': saved.data['revision']}, format='json')
        self.assertEqual(published.status_code, 200)
        self.assertEqual(self.lesson.blocks.get().content, block['content'])

    def test_presentation_reading_order_and_slide_provenance(self):
        job = self.enqueue('pptx'); execute(claim()); job.refresh_from_db()
        self.assertEqual(job.status, 'SUCCEEDED', job.error_code)
        self.assertEqual([x['text'] for x in job.result['elements']], ['Above', 'Below'])
        self.assertEqual([x['page'] for x in job.result['elements']], [1, 1])
        self.assertEqual(job.result['elements'][0]['bbox']['unit'], 'emu')

    def test_validation_active_job_and_cross_lesson_access(self):
        bad = self.client.post(self.url, {'file': SimpleUploadedFile('fake.pdf', b'not a pdf')}, format='multipart')
        self.assertEqual(bad.status_code, 400)
        self.assertEqual(LectureImportJob.objects.count(), 0)
        job = self.enqueue()
        self.assertEqual(self.client.post(self.url, {'file': office_file()}, format='multipart').status_code, 409)
        other = Lesson.objects.create(module=self.lesson.module, title='Other')
        self.assertEqual(self.client.get(f'/api/v1/lessons/{other.short_id}/imports/{job.pk}/').status_code, 404)
        self.assertEqual(self.student.post(self.url, {'file': office_file()}, format='multipart').status_code, 403)

    def test_revision_conflict_retains_saved_draft_and_blocks_stale_publish(self):
        original = self.client.get(self.draft_url).data
        saved = self.client.put(self.draft_url, {**original, 'title': 'First tab'}, format='json')
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(self.client.put(self.draft_url, {**original, 'title': 'Stale tab'}, format='json').status_code, 409)
        self.assertEqual(self.client.post(self.draft_url, {'revision': original['revision']}, format='json').status_code, 409)
        self.lesson.refresh_from_db()
        self.assertEqual(self.lesson.draft_data['title'], 'First tab')
        self.assertEqual(self.lesson.status, 'DRAFT')

    def test_lease_recovery_and_stale_worker_cannot_commit(self):
        self.enqueue(); old = claim()
        self.assertIsNone(claim())
        LectureImportJob.objects.filter(pk=old.pk).update(lease_expires_at=timezone.now() - timedelta(seconds=1))
        fresh = claim()
        self.assertEqual(fresh.attempts, 2)
        with TemporaryDirectory() as folder:
            with self.assertRaises(ImportFailure):
                store_result(old, {'assets': [], 'elements': [], 'warnings': []}, folder)
        LectureImportJob.objects.filter(pk=old.pk).update(lease_expires_at=timezone.now() - timedelta(seconds=1))
        self.assertIsNone(claim())
        fresh.refresh_from_db(); self.assertEqual(fresh.error_code, 'worker_interrupted')

    def test_import_assets_survive_cleanup_and_remain_private(self):
        from PIL import Image
        self.enqueue(); job = claim()
        with TemporaryDirectory() as folder:
            Image.new('RGB', (8, 8), 'white').save(Path(folder) / 'figure.png')
            store_result(job, {'assets': [{'key': 'figure', 'file': 'figure.png', 'is_page': False}],
                              'elements': [{'source_id': 'e1', 'kind': 'image', 'image': 'figure', 'text': ''}], 'warnings': []}, folder)
        job.refresh_from_db(); media_id = job.result['blocks'][0]['media']
        self.assertEqual(self.student.get(f'/api/v1/media/{media_id}/').status_code, 403)
        self.assertEqual(self.client.delete(f'/api/v1/media/{media_id}/').status_code, 400)
        call_command('cleanup_unused_media', delete=True, stdout=io.StringIO())
        image = self.client.get(f'/api/v1/media/{media_id}/')
        self.assertEqual(image.status_code, 200); image.close()

    def test_worker_timeout_is_visible_without_draft_changes(self):
        self.enqueue(); job = claim()
        with patch('apps.lessons.importing.worker.run_child', side_effect=ImportFailure('timeout')):
            execute(job)
        response = self.client.get(self.url + f'{job.pk}/')
        self.assertEqual(response.data['status'], 'FAILED')
        self.assertIn('времени', response.data['error'])
        self.lesson.refresh_from_db(); self.assertIsNone(self.lesson.draft_data)


class StructureTests(SimpleTestCase):
    def test_langchain_only_changes_roles_with_exact_source_coverage(self):
        from apps.lessons.importing.structure import classify
        elements = [{'source_id': 'e1', 'kind': 'paragraph', 'text': 'Source 42 <script> untouched', 'needs_structure_review': True}]
        config = {'base_url': 'http://example.test', 'api_key': '', 'model': 'test', 'timeout': 1}
        with patch('apps.lessons.importing.structure.call_api', return_value={'payload': {'elements': [{'source_id': 'e1', 'kind': 'quote'}]}}):
            self.assertEqual(classify(elements, config), [])
        self.assertEqual(elements[0]['kind'], 'quote')
        self.assertEqual(elements[0]['text'], 'Source 42 <script> untouched')
        elements[0]['kind'] = 'paragraph'
        with patch('apps.lessons.importing.structure.call_api', return_value={'payload': {'elements': [{'source_id': 'invented', 'kind': 'heading'}]}}):
            self.assertEqual(classify(elements, config), ['ai_unavailable'])
        self.assertEqual(elements[0]['kind'], 'paragraph')

    def test_trusted_native_structure_does_not_call_ai(self):
        from apps.lessons.importing.structure import classify
        elements = [{'source_id': 'e1', 'kind': 'heading', 'text': 'Native heading'},
                    {'source_id': 'e2', 'kind': 'paragraph', 'text': 'An ordinary paragraph with no ambiguous formatting.'},
                    {'source_id': 'e3', 'kind': 'quote', 'text': 'An explicit quotation.'}]
        with patch('apps.lessons.importing.structure.call_api') as transport:
            self.assertEqual(classify(elements, {'base_url': 'http://example.test', 'model': 'test'}), [])
        transport.assert_not_called()

    def test_only_ambiguous_elements_are_sent_to_ai(self):
        from apps.lessons.importing.structure import classify
        elements = [{'source_id': 'e1', 'kind': 'heading', 'text': 'Native heading'},
                    {'source_id': 'e2', 'kind': 'paragraph', 'text': 'Normal body paragraph.'},
                    {'source_id': 'e3', 'kind': 'paragraph', 'text': 'POSSIBLE HEADING', 'inline': [{'text': 'POSSIBLE HEADING', 'marks': [{'type': 'bold'}]}]}]
        with patch('apps.lessons.importing.structure.call_api', return_value={'payload': {'elements': [{'source_id': 'e3', 'kind': 'heading', 'level': 3}]}}) as transport:
            self.assertEqual(classify(elements, {'base_url': 'http://example.test', 'api_key': '', 'model': 'test', 'timeout': 1}), [])
        request = json.loads(transport.call_args.kwargs['messages'][1]['content'])
        self.assertEqual([item['source_id'] for item in request], ['e3'])
        self.assertEqual(elements[1]['kind'], 'paragraph')
        self.assertEqual(elements[2]['kind'], 'heading')
        self.assertEqual(elements[2]['text'], 'POSSIBLE HEADING')

    def test_ocr_is_selected_only_for_images_and_damaged_text(self):
        from apps.lessons.importing.quality import pdf_ocr_mode, page_ranges
        self.assertIsNone(pdf_ocr_mode('Digital source text'))
        self.assertIsNone(pdf_ocr_mode('', has_objects=False))
        self.assertEqual(pdf_ocr_mode('Digital text plus screenshot', has_images=True), 'regions')
        self.assertEqual(pdf_ocr_mode('', has_images=True), 'full')
        self.assertEqual(pdf_ocr_mode('\ufffd\ufffd\ufffd broken encoding'), 'full')
        self.assertEqual(page_ranges([2, 3, 5, 8, 9]), [(2, 3), (5, 5), (8, 9)])

    def test_pdf_fallback_combines_native_and_ocr_pages_in_source_order(self):
        from apps.lessons.importing.parser import Extractor
        native = Mock(); native.convert.return_value = SimpleNamespace(status=SimpleNamespace(value='success'), document='native')
        ocr = Mock(); ocr.convert.return_value = SimpleNamespace(status=SimpleNamespace(value='success'), document='ocr')
        with TemporaryDirectory() as folder:
            extractor = Extractor(folder, {'max_pages': 100, 'languages': 'rus+eng+kaz'})
            def consume(document, pages, ocr=False):
                for page in sorted(pages):
                    extractor.add('paragraph', f'{document} page {page}', page=page, ocr=ocr)
            with patch('apps.lessons.importing.parser.pdf_page_modes', return_value={1: None, 2: 'full', 3: None}), \
                    patch('apps.lessons.importing.parser.pdf_converter', side_effect=[native, ocr]) as factory, \
                    patch.object(extractor, 'require_ocr'), patch.object(extractor, 'pdf_items', side_effect=consume):
                result = extractor.extract(Path(folder) / 'mixed.pdf')
            self.assertEqual([entry['text'] for entry in result['elements']], ['native page 1', 'ocr page 2', 'native page 3'])
            self.assertEqual([entry['source_id'] for entry in result['elements']], ['e1', 'e2', 'e3'])
            self.assertEqual(factory.call_args_list[0].args, ('rus+eng+kaz',))
            self.assertEqual(factory.call_args_list[1].args, ('rus+eng+kaz', 'full'))
            ocr.convert.assert_called_once_with(Path(folder) / 'mixed.pdf', page_range=(2, 2))

    def test_failed_native_conversion_automatically_retries_with_full_ocr(self):
        from apps.lessons.importing.parser import Extractor
        native = Mock(); native.convert.side_effect = ValueError('native conversion failed')
        ocr = Mock(); ocr.convert.return_value = SimpleNamespace(status=SimpleNamespace(value='success'), document='ocr')
        with TemporaryDirectory() as folder:
            extractor = Extractor(folder, {'max_pages': 100, 'languages': 'eng'})
            with patch('apps.lessons.importing.parser.pdf_page_modes', return_value={1: None}), \
                    patch('apps.lessons.importing.parser.pdf_converter', side_effect=[native, ocr]), \
                    patch.object(extractor, 'require_ocr'), \
                    patch.object(extractor, 'pdf_items', side_effect=lambda *args, **kwargs: extractor.add('paragraph', 'Recovered source', page=1, ocr=True)):
                result = extractor.extract(Path(folder) / 'digital.pdf')
            self.assertEqual(result['elements'][0]['text'], 'Recovered source')
            self.assertIn('ocr_review', result['warnings'])
            ocr.convert.assert_called_once_with(Path(folder) / 'digital.pdf', page_range=(1, 1))

    def test_real_mixed_pdf_keeps_digital_pages_out_of_ocr(self):
        import pypdfium2 as pdfium
        from PIL import Image
        from apps.lessons.test_import_ocr import make_pdf
        from apps.lessons.importing.parser import pdf_page_modes
        with TemporaryDirectory() as folder:
            digital, scan, mixed = [Path(folder) / filename for filename in ('digital.pdf', 'scan.pdf', 'mixed.pdf')]
            make_pdf(digital)
            make_pdf(scan, Image.new('RGB', (64, 64), 'white'))
            with pdfium.PdfDocument.new() as combined, pdfium.PdfDocument(digital) as first, pdfium.PdfDocument(scan) as second:
                combined.import_pages(first)
                combined.import_pages(second)
                combined.save(mixed)
            self.assertEqual(pdf_page_modes(mixed, max_pages=10), {1: None, 2: 'full'})
            with self.assertRaises(ImportFailure):
                pdf_page_modes(mixed, max_pages=1)

    def test_oversized_element_is_rejected_without_truncation(self):
        with self.assertRaises(ImportFailure):
            to_blocks([{'kind': 'paragraph', 'text': 'x' * 50000}], {})

    def test_ocr_text_is_extracted_next_to_the_image(self):
        from PIL import Image
        from apps.lessons.importing.parser import Extractor
        data = io.BytesIO(); Image.new('RGB', (64, 64), 'white').save(data, 'PNG')
        with TemporaryDirectory() as folder:
            extractor = Extractor(folder, {'languages': 'rus+eng+kaz'})
            with patch('apps.lessons.importing.parser.subprocess.run') as run:
                run.return_value.returncode = 0; run.return_value.stdout = 'Recognized source 42'
                extractor.image(data.getvalue(), page=2)
            self.assertEqual([x['kind'] for x in extractor.elements], ['image', 'paragraph'])
            self.assertEqual(extractor.elements[1]['text'], 'Recognized source 42')
            self.assertEqual(extractor.elements[1]['page'], 2)
            self.assertIn('rus+eng+kaz', run.call_args.args[0])

    def test_word_hyperlink_text_is_not_lost_between_formatted_runs(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from apps.lessons.importing.parser import Extractor
        with TemporaryDirectory() as folder:
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run('Before ').bold = True
            link, run, text = OxmlElement('w:hyperlink'), OxmlElement('w:r'), OxmlElement('w:t')
            text.text = 'linked source'
            run.append(text); link.append(run); paragraph._p.append(link)
            paragraph.add_run(' after')
            path = Path(folder) / 'lecture.docx'; document.save(path)
            result = Extractor(folder, {}).extract(path)
            blocks = to_blocks(result['elements'], {})
            node = json.loads(blocks[0]['content'])['content'][0]
            self.assertEqual(''.join(item['text'] for item in node['content']), 'Before linked source after')
